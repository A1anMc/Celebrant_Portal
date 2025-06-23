import os
import pickle
import logging
import re
from datetime import datetime
from typing import List, Optional, Dict, Any
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from google.auth.transport.requests import Request
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError
import html
from bs4 import BeautifulSoup

# Configure logging
logger = logging.getLogger(__name__)

# Drive API scopes - read-only access to Drive files
DRIVE_SCOPES = [
    'https://www.googleapis.com/auth/gmail.modify',  # Existing Gmail scope
    'https://www.googleapis.com/auth/drive.readonly'  # New Drive scope
]

class DriveService:
    def __init__(self):
        self.service = None
        self.creds = None
        self.SCOPES = DRIVE_SCOPES

    def ensure_authenticated(self) -> None:
        """Ensure we have valid credentials with Drive access."""
        print("Checking for credentials.json...")
        if not os.path.exists('credentials.json'):
            raise Exception(
                "Google API credentials not found. Please follow these steps:\n"
                "1. Go to Google Cloud Console (https://console.cloud.google.com)\n"
                "2. Create a project or select an existing one\n"
                "3. Enable the Gmail API and Google Drive API\n"
                "4. Go to Credentials\n"
                "5. Create OAuth 2.0 Client ID (Application type: Desktop)\n"
                "6. Download the credentials and save as 'credentials.json'\n"
                "7. Place the file in the application root directory"
            )
                
        print("Checking for token.pickle...")
        if os.path.exists('token.pickle'):
            print("Found token.pickle, loading credentials...")
            with open('token.pickle', 'rb') as token:
                self.creds = pickle.load(token)
                
                # Check if we have the Drive scope
                if self.creds and hasattr(self.creds, 'scopes'):
                    if 'https://www.googleapis.com/auth/drive.readonly' not in self.creds.scopes:
                        print("Drive scope not found in existing credentials, need to re-authenticate...")
                        os.remove('token.pickle')
                        self._run_oauth_flow()
                        return
                
                # Check if credentials are expired and refresh if possible
                if self.creds and self.creds.expired and self.creds.refresh_token:
                    print("Credentials expired, attempting to refresh...")
                    try:
                        self.creds.refresh(Request())
                        print("Successfully refreshed credentials")
                        # Save the refreshed credentials
                        with open('token.pickle', 'wb') as token_file:
                            pickle.dump(self.creds, token_file)
                    except Exception as e:
                        print(f"Failed to refresh credentials: {str(e)}")
                        os.remove('token.pickle')
                        self._run_oauth_flow()
                        return
                
                if self.creds and not self.creds.expired:
                    print("Building Drive service with valid credentials...")
                    self.service = build('drive', 'v3', credentials=self.creds)
                else:
                    print("Credentials expired and couldn't be refreshed")
                    os.remove('token.pickle')
                    self._run_oauth_flow()
        else:
            print("No token.pickle found, starting OAuth flow...")
            self._run_oauth_flow()

    def _run_oauth_flow(self):
        """Run the OAuth flow to get new credentials."""
        try:
            flow = InstalledAppFlow.from_client_secrets_file(
                'credentials.json', self.SCOPES)
            self.creds = flow.run_local_server(port=0)
            print("Successfully obtained new credentials with Drive access")
            
            # Save the credentials for the next run
            with open('token.pickle', 'wb') as token:
                pickle.dump(self.creds, token)
            print("Saved credentials to token.pickle")
            
            self.service = build('drive', 'v3', credentials=self.creds)
        except Exception as e:
            print(f"Error during OAuth flow: {str(e)}")
            raise Exception(
                f"Failed to authenticate with Google Drive: {str(e)}. "
                "Please ensure you have valid credentials.json and "
                "proper permissions."
            )

    def search_template_files(self, folder_name: Optional[str] = None) -> List[Dict[str, Any]]:
        """Search for template files in Google Drive supporting multiple formats.
        
        Args:
            folder_name: Optional folder name to search within
            
        Returns:
            List of file metadata dictionaries
        """
        try:
            if not self.service:
                self.ensure_authenticated()

            # Build search query for multiple file types
            query_parts = []
            
            # Search for multiple file types
            mime_types = [
                "mimeType='application/vnd.google-apps.document'",  # Google Docs
                "mimeType='application/vnd.google-apps.spreadsheet'",  # Google Sheets
                "mimeType='application/vnd.openxmlformats-officedocument.wordprocessingml.document'",  # .docx
                "mimeType='application/msword'",  # .doc
                "mimeType='application/pdf'",  # .pdf
                "mimeType='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'",  # .xlsx
                "mimeType='application/vnd.ms-excel'"  # .xls
            ]
            
            mime_query = "(" + " or ".join(mime_types) + ")"
            query_parts.append(mime_query)
            
            # Search for files with template-related keywords in the name
            template_keywords = ['template', 'ceremony', 'mc', 'vow', 'wedding', 'celebrant', 'script', 'order of service']
            keyword_query = " or ".join([f"name contains '{keyword}'" for keyword in template_keywords])
            query_parts.append(f"({keyword_query})")
            
            # Add folder restriction if specified
            if folder_name:
                # First find the folder
                folder_query = f"name='{folder_name}' and mimeType='application/vnd.google-apps.folder'"
                folder_results = self.service.files().list(
                    q=folder_query,
                    fields="files(id, name)"
                ).execute()
                
                folders = folder_results.get('files', [])
                if folders:
                    folder_id = folders[0]['id']
                    query_parts.append(f"'{folder_id}' in parents")
            
            # Only include files that are not trashed
            query_parts.append("trashed=false")
            
            query = " and ".join(query_parts)
            
            print(f"Searching Drive with query: {query}")
            
            # Execute search
            results = self.service.files().list(
                q=query,
                fields="files(id, name, createdTime, modifiedTime, size, owners, mimeType)",
                orderBy="modifiedTime desc",
                pageSize=50
            ).execute()
            
            files = results.get('files', [])
            
            # Format the results with proper type detection
            template_files = []
            for file in files:
                file_type = self._get_file_type(file.get('mimeType', ''))
                template_files.append({
                    'id': file['id'],
                    'name': file['name'],
                    'created_time': file.get('createdTime'),
                    'modified_time': file.get('modifiedTime'),
                    'size': file.get('size'),
                    'owners': file.get('owners', []),
                    'mime_type': file.get('mimeType'),
                    'type': file_type,
                    'can_preview': self._can_preview(file.get('mimeType', '')),
                    'can_import': self._can_import(file.get('mimeType', ''))
                })
            
            print(f"Found {len(template_files)} template files")
            return template_files
            
        except HttpError as error:
            logger.error(f"Drive API error: {error}")
            raise Exception(f"Failed to search Drive files: {error}")
        except Exception as e:
            logger.error(f"Unexpected error searching Drive: {e}")
            raise Exception(f"Failed to search Drive files: {str(e)}")
    
    def _get_file_type(self, mime_type: str) -> str:
        """Get user-friendly file type from MIME type."""
        type_mapping = {
            'application/vnd.google-apps.document': 'google_doc',
            'application/vnd.google-apps.spreadsheet': 'google_sheet',
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document': 'word_docx',
            'application/msword': 'word_doc',
            'application/pdf': 'pdf',
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet': 'excel_xlsx',
            'application/vnd.ms-excel': 'excel_xls'
        }
        return type_mapping.get(mime_type, 'unknown')
    
    def _can_preview(self, mime_type: str) -> bool:
        """Check if file type can be previewed."""
        previewable_types = [
            'application/vnd.google-apps.document',
            'application/vnd.google-apps.spreadsheet',
            'application/pdf'
        ]
        return mime_type in previewable_types
    
    def _can_import(self, mime_type: str) -> bool:
        """Check if file type can be imported as template."""
        importable_types = [
            'application/vnd.google-apps.document',
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            'application/msword',
            'application/pdf'
        ]
        return mime_type in importable_types

    def get_file_content(self, file_id: str) -> Dict[str, Any]:
        """Get the content of a file supporting multiple formats.
        
        Args:
            file_id: The Google Drive file ID
            
        Returns:
            Dictionary with file metadata and content
        """
        try:
            if not self.service:
                self.ensure_authenticated()

            # Get file metadata including MIME type
            file_metadata = self.service.files().get(
                fileId=file_id,
                fields="id, name, createdTime, modifiedTime, size, owners, mimeType"
            ).execute()
            
            mime_type = file_metadata.get('mimeType', '')
            file_type = self._get_file_type(mime_type)
            
            # Handle different file types
            if mime_type == 'application/vnd.google-apps.document':
                # Google Docs - export as HTML
                content = self.service.files().export(
                    fileId=file_id,
                    mimeType='text/html'
                ).execute()
                html_content = content.decode('utf-8')
                
                return {
                    'metadata': file_metadata,
                    'file_type': file_type,
                    'html_content': html_content,
                    'processed_content': self._process_html_content(html_content)
                }
                
            elif mime_type == 'application/vnd.google-apps.spreadsheet':
                # Google Sheets - export as HTML
                content = self.service.files().export(
                    fileId=file_id,
                    mimeType='text/html'
                ).execute()
                html_content = content.decode('utf-8')
                
                return {
                    'metadata': file_metadata,
                    'file_type': file_type,
                    'html_content': html_content,
                    'processed_content': self._process_spreadsheet_content(html_content)
                }
                
            elif mime_type in ['application/vnd.openxmlformats-officedocument.wordprocessingml.document', 'application/msword']:
                # Word documents - download and extract text
                content = self.service.files().get_media(fileId=file_id).execute()
                
                return {
                    'metadata': file_metadata,
                    'file_type': file_type,
                    'raw_content': content,
                    'processed_content': self._process_word_content(content, mime_type)
                }
                
            elif mime_type == 'application/pdf':
                # PDF files - download raw content
                content = self.service.files().get_media(fileId=file_id).execute()
                
                return {
                    'metadata': file_metadata,
                    'file_type': file_type,
                    'raw_content': content,
                    'processed_content': self._process_pdf_content(content)
                }
                
            elif mime_type in ['application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'application/vnd.ms-excel']:
                # Excel files - download and extract text
                content = self.service.files().get_media(fileId=file_id).execute()
                
                return {
                    'metadata': file_metadata,
                    'file_type': file_type,
                    'raw_content': content,
                    'processed_content': self._process_excel_content(content, mime_type)
                }
                
            else:
                raise Exception(f"Unsupported file type: {mime_type}")
            
        except HttpError as error:
            logger.error(f"Drive API error getting file {file_id}: {error}")
            raise Exception(f"Failed to get file content: {error}")
        except Exception as e:
            logger.error(f"Unexpected error getting file content: {e}")
            raise Exception(f"Failed to get file content: {str(e)}")

    def _process_html_content(self, html_content: str) -> str:
        """Process HTML content to create a template.
        
        Args:
            html_content: Raw HTML content from Google Docs
            
        Returns:
            Processed template content with placeholders
        """
        try:
            # Parse HTML
            soup = BeautifulSoup(html_content, 'html.parser')
            
            # Remove style tags and script tags
            for tag in soup(['style', 'script', 'meta', 'link']):
                tag.decompose()
            
            # Get the body content or full content if no body
            body = soup.find('body')
            if body:
                content = str(body)
            else:
                content = str(soup)
            
            # Clean up the HTML
            content = self._clean_html(content)
            
            # Replace common names with placeholders
            content = self._replace_names_with_placeholders(content)
            
            return content
            
        except Exception as e:
            logger.error(f"Error processing HTML content: {e}")
            # Fallback to plain text if HTML processing fails
            return self._html_to_text(html_content)

    def _clean_html(self, html_content: str) -> str:
        """Clean HTML content by removing unnecessary attributes and tags."""
        soup = BeautifulSoup(html_content, 'html.parser')
        
        # Remove unnecessary attributes but keep basic formatting
        allowed_tags = ['p', 'div', 'span', 'br', 'strong', 'b', 'em', 'i', 'u', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'ul', 'ol', 'li']
        allowed_attrs = ['style']  # Keep minimal styling
        
        for tag in soup.find_all(True):
            if tag.name not in allowed_tags:
                tag.unwrap()  # Remove tag but keep content
            else:
                # Clean attributes
                attrs_to_remove = []
                for attr in tag.attrs:
                    if attr not in allowed_attrs:
                        attrs_to_remove.append(attr)
                for attr in attrs_to_remove:
                    del tag[attr]
        
        return str(soup)

    def _replace_names_with_placeholders(self, content: str) -> str:
        """Replace common names and wedding terms with template placeholders."""
        
        # Common placeholder replacements
        replacements = {
            # Names - these would typically be replaced with actual names from the document
            r'\b[A-Z][a-z]+ and [A-Z][a-z]+\b': '{{ partner1_name }} and {{ partner2_name }}',
            r'\b[A-Z][a-z]+\s+&\s+[A-Z][a-z]+\b': '{{ partner1_name }} & {{ partner2_name }}',
            
            # Common wedding terms
            r'\bthe bride\b': '{{ partner1_name }}',
            r'\bthe groom\b': '{{ partner2_name }}',
            r'\bThe Bride\b': '{{ partner1_name }}',
            r'\bThe Groom\b': '{{ partner2_name }}',
            
            # Date placeholders
            r'\b\d{1,2}(st|nd|rd|th)?\s+(January|February|March|April|May|June|July|August|September|October|November|December)\s+\d{4}\b': '{{ ceremony_date }}',
            r'\b(January|February|March|April|May|June|July|August|September|October|November|December)\s+\d{1,2}(st|nd|rd|th)?,?\s+\d{4}\b': '{{ ceremony_date }}',
            
            # Time placeholders
            r'\b\d{1,2}:\d{2}\s*(am|pm|AM|PM)\b': '{{ ceremony_time }}',
            
            # Location placeholders
            r'\bat\s+[A-Z][^.!?]*(?=\.|!|\?|$)': 'at {{ ceremony_location }}',
        }
        
        # Apply replacements
        for pattern, replacement in replacements.items():
            content = re.sub(pattern, replacement, content, flags=re.IGNORECASE)
        
        return content

    def _html_to_text(self, html_content: str) -> str:
        """Convert HTML to plain text as fallback."""
        try:
            soup = BeautifulSoup(html_content, 'html.parser')
            text = soup.get_text()
            
            # Clean up whitespace
            lines = [line.strip() for line in text.splitlines()]
            text = '\n'.join(line for line in lines if line)
            
            return text
        except Exception as e:
            logger.error(f"Error converting HTML to text: {e}")
            return html_content

    def preview_file(self, file_id: str) -> Dict[str, Any]:
        """Get a preview of a file without full processing.
        
        Args:
            file_id: The Google Drive file ID
            
        Returns:
            Dictionary with preview information
        """
        try:
            if not self.service:
                self.ensure_authenticated()

            # Get file metadata
            file_metadata = self.service.files().get(
                fileId=file_id,
                fields="id, name, createdTime, modifiedTime, size, owners, description"
            ).execute()
            
            # Get first 1000 characters as plain text preview
            try:
                content = self.service.files().export(
                    fileId=file_id,
                    mimeType='text/plain'
                ).execute()
                
                preview_text = content.decode('utf-8')[:1000]
                if len(content.decode('utf-8')) > 1000:
                    preview_text += "..."
                    
            except Exception:
                # Fallback to HTML export if plain text fails
                content = self.service.files().export(
                    fileId=file_id,
                    mimeType='text/html'
                ).execute()
                
                html_content = content.decode('utf-8')
                soup = BeautifulSoup(html_content, 'html.parser')
                preview_text = soup.get_text()[:1000]
                if len(soup.get_text()) > 1000:
                    preview_text += "..."
            
            return {
                'metadata': file_metadata,
                'preview': preview_text
            }
            
        except HttpError as error:
            logger.error(f"Drive API error previewing file {file_id}: {error}")
            raise Exception(f"Failed to preview file: {error}")
        except Exception as e:
            logger.error(f"Unexpected error previewing file: {e}")
            raise Exception(f"Failed to preview file: {str(e)}")

    def check_drive_access(self) -> bool:
        """Check if we have valid Drive API access."""
        try:
            if not self.service:
                self.ensure_authenticated()
            
            # Try a simple API call
            self.service.files().list(pageSize=1).execute()
            return True
            
        except Exception as e:
            logger.error(f"Drive access check failed: {e}")
            return False
    
    def _process_spreadsheet_content(self, html_content: str) -> str:
        """Process Google Sheets HTML content."""
        try:
            soup = BeautifulSoup(html_content, 'html.parser')
            
            # Find tables and convert to readable format
            tables = soup.find_all('table')
            processed_content = ""
            
            for i, table in enumerate(tables):
                if i > 0:
                    processed_content += "\n\n"
                
                processed_content += f"Table {i+1}:\n"
                processed_content += "=" * 20 + "\n"
                
                rows = table.find_all('tr')
                for row in rows:
                    cells = row.find_all(['td', 'th'])
                    row_text = " | ".join(cell.get_text().strip() for cell in cells)
                    if row_text.strip():
                        processed_content += row_text + "\n"
            
            return processed_content or "Could not extract content from spreadsheet"
            
        except Exception as e:
            logger.error(f"Error processing spreadsheet content: {e}")
            return "Error processing spreadsheet content"
    
    def _process_word_content(self, content: bytes, mime_type: str) -> str:
        """Process Word document content."""
        try:
            import io
            from docx import Document
            
            if mime_type == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document':
                # .docx files
                doc = Document(io.BytesIO(content))
                text_content = []
                
                for paragraph in doc.paragraphs:
                    if paragraph.text.strip():
                        text_content.append(paragraph.text)
                
                processed_content = '\n\n'.join(text_content)
                return self._replace_names_with_placeholders(processed_content)
            
            else:
                # .doc files - more complex, fallback to basic text extraction
                # This would require additional libraries like python-docx2txt or antiword
                return "Word .doc files require additional processing. Please convert to .docx or Google Docs format."
                
        except ImportError:
            return "python-docx library required for Word document processing. Please install it or convert to Google Docs format."
        except Exception as e:
            logger.error(f"Error processing Word content: {e}")
            return f"Error processing Word document: {str(e)}"
    
    def _process_pdf_content(self, content: bytes) -> str:
        """Process PDF content."""
        try:
            import PyPDF2
            import io
            
            pdf_reader = PyPDF2.PdfReader(io.BytesIO(content))
            text_content = []
            
            for page in pdf_reader.pages:
                page_text = page.extract_text()
                if page_text.strip():
                    text_content.append(page_text)
            
            processed_content = '\n\n'.join(text_content)
            return self._replace_names_with_placeholders(processed_content)
            
        except ImportError:
            return "PyPDF2 library required for PDF processing. Please install it or convert to Google Docs format."
        except Exception as e:
            logger.error(f"Error processing PDF content: {e}")
            return f"Error processing PDF: {str(e)}"
    
    def _process_excel_content(self, content: bytes, mime_type: str) -> str:
        """Process Excel content."""
        try:
            import pandas as pd
            import io
            
            # Read Excel file
            if mime_type == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet':
                df = pd.read_excel(io.BytesIO(content), engine='openpyxl')
            else:
                df = pd.read_excel(io.BytesIO(content), engine='xlrd')
            
            # Convert to readable format
            processed_content = "Excel Content:\n"
            processed_content += "=" * 20 + "\n"
            processed_content += df.to_string(index=False)
            
            return processed_content
            
        except ImportError:
            return "pandas and openpyxl/xlrd libraries required for Excel processing. Please install them or convert to Google Sheets format."
        except Exception as e:
            logger.error(f"Error processing Excel content: {e}")
            return f"Error processing Excel file: {str(e)}"